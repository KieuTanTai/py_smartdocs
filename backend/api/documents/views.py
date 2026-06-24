from __future__ import annotations

import hashlib
import json
import os
import time
from pathlib import Path
import traceback

from backend.apps.config import container
from backend.apps.core.interfaces.system.i_config import IConfigProvider
from backend.apps.core.interfaces.system.i_logging import ILogger
import numpy as np
import pypdf
from django.conf import settings
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework import status

from backend.apps.services.chat.models import DocumentModel
from backend.apps.core.normalize.normalize import Normalize
from backend.apps.core.chunk.chunker import Chunker
from backend.apps.llm.llm_provider_factory import LLMProviderFactory
from backend.apps.core.interfaces.dataclass.i_dataclass_transaction import ICompletionRequest
from backend.apps.core.enums.e_provider_name import EProviderName
from backend.apps.core.enums.e_backend_storage_name import EBackendStorageName
from backend.apps.services.rag_base.locate.locate_service import LocateService

from backend.apps.application.conversations.application import ConversationApplication
from sys_services.system_dirs import METADATA_DIR

# Singleton application instances
__container = container.BackendContainer()


class DocumentListView(APIView):
    def get(self, request):
        doc_app = __container.document_application()
        
        result = doc_app.list_files(request.GET.get("conversation_id", ""), file_caller="DocumentListView")
        data = []
        for file in result:
            data.append({
                "id": str(file.conversation_files_id),
                "cloud_id": str(file.conversation_files_cloud_id),
                "vector_file": str(file.conversation_files_document.documents_file_path),
                "uploaded_at": str(file.conversation_files_uploaded_at),
            })
        return Response(data, status=status.HTTP_200_OK)

class DocumentUploadView(APIView):
    def post(self, request):
        sys_logger = __container.log_pool() # Lấy ILogger từ Container

        uploaded_file = request.FILES.get("file")
        if not uploaded_file:
            sys_logger.warning("No file uploaded in request", source="DocumentUploadView", call_by="post")
            return Response({"error": "No file uploaded"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            doc_app = __container.document_application()
            file_content = uploaded_file.read()
            
            doc_request = doc_app.upload_document(
                request=request,
                file_caller = Path(__file__).stem,
            )
            
            sys_logger.info(f"File uploaded successfully: {uploaded_file.name}", source="DocumentUploadView", call_by="post")
            return Response({
                "title": uploaded_file.name,
                "status": "uploaded"
            }, status=status.HTTP_201_CREATED)
            
        except ValueError as e:
            sys_logger.error(f"Validation Error: {e}", source="DocumentUploadView", call_by="post", method_call="upload_document")
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as e:
            sys_logger.error(f"Upload failed: {e}\n{traceback.format_exc()}", source="DocumentUploadView", call_by="post")
            return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


class DocumentDetailView(APIView):
    def get(self, request, document_id: str):
        try:
            doc_app = __container.document_application()
            doc = doc_app.get_document(document_id)
            return Response({
                "id": str(doc["id"]),
                "title": doc.get("title", ""),
                "status": doc.get("status", "unknown")
            }, status=status.HTTP_200_OK)
        except ValueError:
            return Response({"error": "Document not found"}, status=status.HTTP_404_NOT_FOUND)

    def delete(self, request, document_id: str):
        sys_logger = __container.log_pool()

        try:
            doc_app = __container.document_application()
            doc_app.delete_document(document_id=document_id, delete_file=True)
            
            try:
                # Xử lý dọn dẹp FAISS
                locate_service = LocateService(metadata_dir=METADATA_DIR, logger=sys_logger)
                faiss_service = locate_service.get_vector_store(EBackendStorageName.FAISS)
                if faiss_service.is_existed_in_metadata(document_id):
                    faiss_service.delete(document_id)
            except Exception as e:
                sys_logger.warning(f"Could not clean up FAISS for {document_id}: {e}", source="DocumentDetailView", call_by="delete")
                
            return Response({"status": "deleted"}, status=status.HTTP_200_OK)
        except ValueError:
            return Response({"error": "Document not found"}, status=status.HTTP_404_NOT_FOUND)
        except Exception as e:
            sys_logger.error(f"Delete failed: {e}", source="DocumentDetailView", call_by="delete")
            return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

class DocumentStatusView(APIView):
    def get(self, request, document_id: str):
        try:
            doc_app = __container.document_application()
            doc = doc_app.get_document(document_id)
            return Response({
                "id": str(doc["id"]),
                "status": doc.get("status", "unknown")
            }, status=status.HTTP_200_OK)
        except ValueError:
            return Response({"error": "Document not found"}, status=status.HTTP_404_NOT_FOUND)


class DocumentIndexView(APIView):
    """
    Index a document: extract text → normalize → chunk → embed → save FAISS index.

    This replaces the old keyword-paragraph approach with a proper vector RAG pipeline:
    1. Extract text from file (PDF/DOCX/TXT)
    2. Normalize whitespace
    3. Chunk into segments (NLTK splitter)
    4. Embed each chunk using Gemini embedding model
    5. Build FAISS index (IndexFlatL2)
    6. Save chunk metadata (id → text) for retrieval
    7. Update DocumentModel.status = "indexed"
    """

    INDEX_EMBEDDING_PROVIDER = EProviderName.GEMINI
    INDEX_EMBEDDING_MODEL = "gemini-embedding-2"

    def post(self, request, document_id: str):
        try:
            doc = DocumentModel.objects.get(pk=document_id)
        except DocumentModel.DoesNotExist:
            return Response({"error": "Document not found"}, status=status.HTTP_404_NOT_FOUND)

        doc.status = "processing"
        doc.save()

        try:
            # 1. Extract text
            extracted_text = self._extract_text(doc)
            if not extracted_text or not extracted_text.strip():
                raise ValueError("No text could be extracted from the document.")

            # 2. Normalize
            normalizer = Normalize()
            normalized_text = normalizer.normalize(extracted_text)

            # 3. Chunk
            chunker = Chunker(logger=ILogger)
            chunks = chunker.create_chunks(normalized_text)
            if not chunks:
                raise ValueError("No chunks generated from document.")

            # 4. Embed chunks
            embeddings = self._embed_chunks(chunks)
            vectors = np.array(embeddings, dtype=np.float32)

            # 5. Build FAISS index and save
            vector_id = str(doc.faiss_index_id)
            locate_service = LocateService(metadata_dir=METADATA_DIR, logger=ILogger)
            faiss_service = locate_service.get_vector_store(EBackendStorageName.FAISS)

            index = faiss_service.create_index(vectors)
            faiss_service.upsert(index=index, conversation_id=doc.faiss_index_id)

            # 6. Save chunk metadata for retrieval (map index → text)
            chunk_metadata = {
                str(idx): chunk_text for idx, chunk_text in enumerate(chunks)
            }
            metadata_path = METADATA_DIR / "docs" / f"{vector_id}.json"
            metadata_path.parent.mkdir(parents=True, exist_ok=True)
            with open(metadata_path, "w", encoding="utf-8") as f:
                json.dump({
                    "document_id": vector_id,
                    "file_name": doc.faiss_index_file_name,
                    "chunk_count": len(chunks),
                    "embedding_model": self.INDEX_EMBEDDING_MODEL,
                    "embedding_dimensions": vectors.shape[1],
                    "chunks": chunk_metadata,
                }, f, ensure_ascii=False, indent=2)

            # 7. Update DB
            doc.content = normalized_text
            doc.status = "indexed"
            doc.save()

            ILogger.info(
                f"Document indexed: {doc.faiss_index_file_name} "
                f"({len(chunks)} chunks, dim={vectors.shape[1]})",
                source="DocumentIndexView",
            )

            return Response({
                "id": str(doc.faiss_index_id),
                "status": "indexed",
                "chunks": len(chunks),
                "dimensions": vectors.shape[1],
            }, status=status.HTTP_200_OK)

        except Exception as e:
            ILogger.error(f"Document indexing failed: {e}", source="DocumentIndexView")
            doc.status = "failed"
            doc.save()
            return Response({"error": str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    def _extract_text(self, doc: DocumentModel) -> str:
        """Extract text from PDF, DOCX, or plain text file."""
        if not doc.file_path or not os.path.exists(doc.file_path):
            return ""

        file_path = doc.file_path

        # Primary method: Use Mistral OCR (via DI Container's extract_content_service)
        try:
            from backend.apps.config.container import BackendContainer
            container = BackendContainer()
            extract_service = container.extract_content_service()
            ocr_response = extract_service.extract(Path(file_path), EProviderName.MISTRAL)
            if ocr_response.extracted_text and ocr_response.extracted_text.strip():
                DEFAULT_LOGGER.info(
                    f"OCR successfully extracted text from: {Path(file_path).name}",
                    source="DocumentIndexView"
                )
                return ocr_response.extracted_text
        except Exception as exc:
            DEFAULT_LOGGER.warning(
                f"OCR extraction failed, falling back to native parser. Error: {exc}",
                source="DocumentIndexView"
            )

        if file_path.endswith(".pdf"):
            try:
                reader = pypdf.PdfReader(file_path)
                text_list = []
                for page in reader.pages:
                    t = page.extract_text()
                    if t:
                        text_list.append(t)
                return "\n".join(text_list)
            except Exception as e:
                ILogger.error(f"Failed to read PDF: {e}", source="DocumentIndexView")
                return ""

        elif file_path.endswith(".docx") or file_path.endswith(".doc"):
            # Try using python-docx if available
            try:
                from docx import Document
                doc_obj = Document(file_path)
                paragraphs = []
                for paragraph in doc_obj.paragraphs:
                    if paragraph.text.strip():
                        paragraphs.append(paragraph.text)
                # Also extract text from tables
                for table in doc_obj.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                paragraphs.append(cell.text)
                extracted = "\n".join(paragraphs)
                if extracted.strip():
                    DEFAULT_LOGGER.info(
                        f"Successfully extracted {len(extracted)} chars from DOCX using python-docx",
                        source="DocumentIndexView"
                    )
                    return extracted
            except ImportError:
                DEFAULT_LOGGER.warning(
                    "python-docx not installed, falling back to XML parsing",
                    source="DocumentIndexView"
                )
            except Exception as e:
                DEFAULT_LOGGER.warning(
                    f"python-docx extraction failed: {e}, falling back to XML parsing",
                    source="DocumentIndexView"
                )
            
            # Fallback to XML parsing
            try:
                import zipfile
                import xml.etree.ElementTree as ET
                with zipfile.ZipFile(file_path) as docx:
                    xml_content = docx.read("word/document.xml")
                    root = ET.fromstring(xml_content)
                    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                    paragraphs = []
                    for p in root.findall(".//w:p", ns):
                        text_runs = [t.text for t in p.findall(".//w:t", ns) if t.text]
                        if text_runs:
                            paragraphs.append("".join(text_runs))
                    extracted = "\n".join(paragraphs)
                    if extracted.strip():
                        DEFAULT_LOGGER.info(
                            f"Successfully extracted {len(extracted)} chars from DOCX using XML parsing",
                            source="DocumentIndexView"
                        )
                        return extracted
                    else:
                        DEFAULT_LOGGER.warning(
                            f"XML parsing returned empty content for DOCX file: {file_path}",
                            source="DocumentIndexView"
                        )
                        return ""
            except Exception as e:
                ILogger.error(f"Failed to read DOCX: {e}", source="DocumentIndexView")
                return ""

        else:
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()
            except Exception as e:
                ILogger.error(f"Failed to read text file: {e}", source="DocumentIndexView")
                return ""

    def _embed_chunks(self, chunks: list[str]) -> list[list[float]]:
        """Embed all chunks using the configured embedding provider."""
        factory = LLMProviderFactory(IConfigProvider, ILogger)
        client = factory.get_provider(self.INDEX_EMBEDDING_PROVIDER)
        embeddings = []

        for chunk in chunks:
            req = ICompletionRequest(
                provider=self.INDEX_EMBEDDING_PROVIDER,
                model=self.INDEX_EMBEDDING_MODEL,
                prompt=chunk,
            )
            resp = client.embedding(req)
            embeddings.append(resp.embedding.tolist())

        return embeddings


class DocumentBulkIndexView(APIView):
    def post(self, request):
        return Response({"detail": "Not implemented."}, status=status.HTTP_501_NOT_IMPLEMENTED)
