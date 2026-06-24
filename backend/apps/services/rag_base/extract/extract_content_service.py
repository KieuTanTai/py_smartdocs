from pathlib import Path
from backend.apps.core.interfaces.dataclass.extract.i_extract_response import IExtractResponse
from backend.apps.core.interfaces.services.rag_base.storage.i_storage import (
    IFileStorage,
)
from backend.apps.core.interfaces.llm.llm_ocr.i_llm_ocr_factory import ILLMOCRFactory
from backend.apps.core.enums.e_provider_name import EProviderName
from backend.apps.core.interfaces.system.i_logging import ILogger
from backend.apps.core.interfaces.services.rag_base.extract.i_extract_content import (
    IExtractContent,
)
from mistralai.client.models import OCRResponse

from backend.apps.utils.is_content_empty import check_empty_content

class ExtractContentService(IExtractContent):

    def __init__(
        self,
        factory: ILLMOCRFactory,
        storage: IFileStorage,
        logger: ILogger,
    ):
        self.factory = factory
        self.storage = storage
        self.logger = logger

    def extract(
        self,
        file_path: Path,
        provider: EProviderName,
        call_by: str = "",
    ) -> IExtractResponse:
        source_log = f"{Path(__file__).parent.absolute()}/{Path(__file__).name}"
        if provider is None:
            raise ValueError("Provider must be specified for extract_from_file_text")
        uploaded_file = self.storage.save_file(file_path)
        # Always use Mistral for OCR since it's currently the only supported OCR provider
        ocr_provider = EProviderName.MISTRAL
        self.logger.info(
            f"Using {ocr_provider.value} for OCR extraction (requested provider: {provider.value})",
            source=source_log,
            call_by=call_by,
            method_call=self.extract.__name__,
        )
        
        try:
            ocr_extractor = self.factory.create_ocr_extractor(ocr_provider)
            ocr_response = ocr_extractor.process_ocr(uploaded_file)
            extracted_text = self.__process_ocr_response(ocr_response, source_log, call_by=call_by)
            return IExtractResponse(uploaded_file.id, extracted_text, ocr_response.model, ocr_response.usage_info.pages_processed, ocr_response.usage_info.doc_size_bytes)
        except (ConnectionError, OSError) as network_error:
            # Network error - fallback to local extraction
            self.logger.warning(
                f"Mistral OCR failed due to network error: {network_error}. Falling back to local extraction.",
                source=source_log,
                call_by=call_by,
                method_call=self.extract.__name__,
            )
            return self.__fallback_local_extraction(file_path, uploaded_file.id, source_log, call_by)
        except Exception as e:
            # Other errors - also try fallback
            self.logger.warning(
                f"Mistral OCR failed: {e}. Attempting fallback to local extraction.",
                source=source_log,
                call_by=call_by,
                method_call=self.extract.__name__,
            )
            return self.__fallback_local_extraction(file_path, uploaded_file.id, source_log, call_by)
    
    def __fallback_local_extraction(
        self,
        file_path: Path,
        file_id: str,
        source_log: str,
        call_by: str = ""
    ) -> IExtractResponse:
        """Fallback to local text extraction when Mistral OCR is unavailable."""
        self.logger.info(
            f"Using local extraction for file: {file_path.name}",
            source=source_log,
            call_by=call_by,
            method_call=self.__fallback_local_extraction.__name__,
        )
        
        extracted_text = ""
        file_ext = file_path.suffix.lower()
        
        try:
            if file_ext == ".pdf":
                # PDF extraction using pypdf
                import pypdf
                with open(file_path, "rb") as f:
                    reader = pypdf.PdfReader(f)
                    text_parts = []
                    for page in reader.pages:
                        text = page.extract_text()
                        if text:
                            text_parts.append(text)
                    extracted_text = "\n".join(text_parts)
                    
            elif file_ext in [".docx", ".doc"]:
                # DOCX extraction using python-docx
                try:
                    from docx import Document
                    doc = Document(file_path)
                    paragraphs = []
                    for paragraph in doc.paragraphs:
                        if paragraph.text.strip():
                            paragraphs.append(paragraph.text)
                    # Also extract from tables
                    for table in doc.tables:
                        for row in table.rows:
                            for cell in row.cells:
                                if cell.text.strip():
                                    paragraphs.append(cell.text)
                    extracted_text = "\n".join(paragraphs)
                except ImportError:
                    # Fallback to XML parsing if python-docx not available
                    import zipfile
                    import xml.etree.ElementTree as ET
                    with zipfile.ZipFile(file_path) as docx:
                        xml_content = docx.read("word/document.xml")
                        root = ET.fromstring(xml_content)
                        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                        paragraphs = []
                        for p in root.findall(".//w:p", ns):
                            text_runs = [t.text for t in p.findall(".//w:t", ns) if t.text]
                            if text_runs:
                                paragraphs.append("".join(text_runs))
                        extracted_text = "\n".join(paragraphs)
                        
            elif file_ext == ".txt":
                # Plain text extraction
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    extracted_text = f.read()
                    
            else:
                # Try to read as text
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    extracted_text = f.read()
            
            if not extracted_text or not extracted_text.strip():
                raise ValueError(f"No text could be extracted from {file_path.name}")
            
            self.logger.info(
                f"Successfully extracted {len(extracted_text)} characters using local extraction",
                source=source_log,
                call_by=call_by,
                method_call=self.__fallback_local_extraction.__name__,
            )
            
            # Return with default values for OCR-specific fields
            return IExtractResponse(
                document_id=file_id,
                extracted_text=extracted_text,
                model="local-extraction",
                pages_processed=1,
                doc_size_bytes=file_path.stat().st_size
            )
            
        except Exception as e:
            error_msg = f"Local extraction also failed for {file_path.name}: {e}"
            self.logger.error(
                error_msg,
                source=source_log,
                call_by=call_by,
                method_call=self.__fallback_local_extraction.__name__,
            )
            raise ValueError(error_msg) from e

    def __process_ocr_response(self, ocr_response: OCRResponse, source_log: str, call_by: str = "") -> str:
        extracted_text = "\n".join([page.markdown for page in ocr_response.pages])
        if self.__validate_response_text(extracted_text, source_log, call_by) is False:
            self.logger.warning(
                f"Extracted text is empty or contains only whitespace. Source log: {source_log}",
                source=str(self.__class__), call_by=call_by, method_call=self.__process_ocr_response.__name__
            )
            raise ValueError("Extracted text is empty or contains only whitespace.")
        return extracted_text

    def __validate_response_text(
        self, text: str, source_log: str, call_by: str = ""
    ) -> bool:
        if check_empty_content(text, source_log, self.logger) is False:
            self.logger.warning(
                f"Extracted text is empty or contains only whitespace. Source log: {source_log}",
                source=str(self.__class__),
                call_by=call_by,
                method_call=self.__validate_response_text.__name__,
            )
            return False
        return True
